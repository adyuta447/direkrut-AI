from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = "output/docx/Template_Lampiran_Proposal_Direkrut_AI.docx"

NAVY = "14324A"
BLUE = "1F6E8C"
TEAL = "2E8A99"
LIGHT_BLUE = "EAF4F7"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "D8DEE3"
DARK_GRAY = "48545C"
WHITE = "FFFFFF"
AMBER = "FFF4D6"
GREEN = "EAF6EE"
RED = "FDECEC"

PAGE_WIDTH_DXA = 11906  # A4
MARGIN_DXA = 1080       # 0.75 inch
CONTENT_DXA = PAGE_WIDTH_DXA - (2 * MARGIN_DXA)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Aptos", size=10.5, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_font(run, size=8.5, color=DARK_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_para(doc, text="", style=None, bold=False, color=None, size=None,
             align=None, before=0, after=6, italic=False, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_font(r, size=size or 10.5, color=color, bold=bold, italic=italic)
    return p


def add_rich_para(doc, parts, before=0, after=6, align=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    for text, kwargs in parts:
        r = p.add_run(text)
        set_font(r, **kwargs)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_font(r, size=10.2)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_font(r, size=10.2)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA - 120])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(f"{label}: ")
    set_font(r, size=10, bold=True, color=NAVY)
    r = p.add_run(text)
    set_font(r, size=10, color=NAVY)
    add_para(doc, "", after=4)


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(header)
        set_font(r, size=font_size, bold=True, color=WHITE)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if row_idx % 2 == 1:
                set_cell_shading(cells[i], LIGHT_GRAY)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            r = p.add_run(str(value))
            set_font(r, size=font_size, color="1F2529")
    set_table_geometry(table, widths)
    add_para(doc, "", after=3)
    return table


def add_attachment_heading(doc, code, title, objective, criteria):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"LAMPIRAN {code}")
    set_font(r, size=10, bold=True, color=TEAL)
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title)
    set_font(r, size=18, bold=True, color=NAVY)
    add_callout(doc, "Tujuan", objective, LIGHT_BLUE)
    add_rich_para(doc, [
        ("Mendukung kriteria: ", {"size": 9.5, "bold": True, "color": DARK_GRAY}),
        (criteria, {"size": 9.5, "color": DARK_GRAY}),
    ], after=10)


def add_template_fields(doc, fields):
    for label, hint in fields:
        add_rich_para(doc, [
            (f"{label}\n", {"size": 9.5, "bold": True, "color": NAVY}),
            (f"[ISI: {hint}]", {"size": 9.5, "italic": True, "color": DARK_GRAY}),
        ], after=8, keep=False)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for style_name, size, color, before, after in [
    ("Heading 1", 18, NAVY, 16, 8),
    ("Heading 2", 13.5, BLUE, 12, 5),
    ("Heading 3", 11.5, NAVY, 8, 4),
]:
    st = styles[style_name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    st = styles[style_name]
    st.font.name = "Aptos"
    st.font.size = Pt(10.2)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.08

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
run = hp.add_run("DIREKRUT AI  /  EVIDENCE PACK")
set_font(run, size=8.5, bold=True, color=TEAL)
footer = section.footer
fp = footer.paragraphs[0]
add_page_field(fp)

# Cover
add_para(doc, "DIREKRUT AI", bold=True, color=TEAL, size=11, after=18)
add_para(doc, "TEMPLATE LAMPIRAN", bold=True, color=NAVY, size=28, after=3)
add_para(doc, "Proposal Digdaya X Hackathon", color=BLUE, size=17, after=18)
add_para(doc, "Evidence Pack untuk Use Case, Feasibility, Complexity, Algorithm & UX, Team, dan Business Plan",
         color=DARK_GRAY, size=11.5, after=28)
add_callout(
    doc,
    "Cara menggunakan",
    "Duplikasi file ini. Ganti seluruh teks [ISI], tempel bukti pada tempat yang ditandai, "
    "hapus bagian yang tidak relevan, lalu ekspor ke PDF. Klaim dalam form harus memakai kode bukti yang sama dengan indeks lampiran.",
    AMBER,
)
add_table(doc,
          ["Metadata", "Isi"],
          [
              ["Team ID", "S0293"],
              ["Nama Tim", "Team Direkrut AI"],
              ["Judul Solusi", "Direkrut AI: Ekosistem Rekrutmen End-to-End Berbasis AI"],
              ["Versi Lampiran", "[ISI: v1.0 / tanggal pembaruan]"],
              ["PIC Evidence Pack", "[ISI: nama dan kontak]"],
              ["Status Kerahasiaan", "Untuk penilaian hackathon - data sensitif telah disamarkan"],
          ],
          [2200, CONTENT_DXA - 2200])
add_para(doc, "Prinsip bukti", style="Heading 2")
for t in [
    "Satu klaim penting = satu kode bukti yang dapat ditelusuri.",
    "Tampilkan artefak aktual, tanggal, metode, dan hasil; jangan hanya menulis testimoni.",
    "Pisahkan status sudah berfungsi, simulasi, dikembangkan, dan direncanakan.",
    "Samarkan NIK, NPWP, alamat pribadi, tanda tangan, nomor dokumen, dan identitas responden yang tidak memberi izin.",
]:
    add_bullet(doc, t)

# Executive guide
doc.add_page_break()
add_para(doc, "PANDUAN EKSEKUTIF", bold=True, color=TEAL, size=10, after=3)
add_para(doc, "Apa yang perlu dilampirkan?", style="Heading 1")
add_para(doc,
         "Untuk proposal saat ini, paket lampiran sebaiknya mengutamakan bukti validasi, bukti produk yang berjalan, "
         "keandalan scoring, kesiapan pilot, legalitas, dan dasar angka finansial. Pitch deck berfungsi sebagai ringkasan, "
         "bukan pengganti bukti rinci.", after=10)
add_table(doc,
          ["Prioritas", "Lampiran", "Kenapa penting", "Status awal"],
          [
              ["P0", "A. Indeks bukti & cross-reference", "Membuat panelis cepat menelusuri klaim", "[ ]"],
              ["P0", "B. Validasi 4 industri + 2 akademisi", "Menguatkan problem, UX, feasibility, dan iterasi", "[ ]"],
              ["P0", "C-D. Use case, UX test, dan change log", "Membuktikan masalah -> fitur -> outcome", "[ ]"],
              ["P0", "E-F. Bukti teknis dan algoritma", "Membuktikan Level 3 dan transparansi keputusan", "[ ]"],
              ["P0", "G-H. MVP, deployment, legal & privasi", "Menutup risiko implementasi dan adopsi", "[ ]"],
              ["P0", "I. Tim dan badan hukum", "Membuktikan ownership dan kesiapan startup", "[ ]"],
              ["P0", "J. Financial analysis & ROI", "Memvalidasi pricing, biaya, unit economics, BEP", "[ ]"],
              ["P1", "K. Pilot, traction, dan growth", "Menunjukkan akses pasar dan calon offtaker", "[ ]"],
              ["P1", "L. Pitch deck", "Ringkasan narasi dan demo untuk panelis", "[ ]"],
          ],
          [850, 2450, 3700, CONTENT_DXA - 7000], font_size=8.6)
add_callout(doc, "Rekomendasi kemasan",
            "Jika portal membatasi ukuran atau halaman, prioritaskan 15-20 halaman evidence pack inti. "
            "Simpan transkrip lengkap, dokumen legal penuh, log pengujian, dan financial model detail sebagai file terpisah atau tautan read-only.",
            GREEN)
add_para(doc, "Catatan penting tentang enam reviewer", style="Heading 2")
add_para(doc,
         "Jangan menyatukan semua masukan menjadi satu paragraf. Untuk setiap reviewer, catat profil, tanggal, metode, "
         "skenario/tugas, kutipan atau temuan, implikasi, keputusan tim, dan bukti perubahan. Kelompokkan peran sebagai "
         "praktisi HR/operations, engineer, product/design, business/industry, serta akademisi yang relevan.", after=8)

# A index
add_attachment_heading(
    doc, "A", "Indeks Bukti dan Cross-Reference",
    "Menghubungkan setiap klaim pada form dengan artefak yang membuktikannya.",
    "Seluruh kriteria penilaian."
)
add_table(doc,
          ["ID", "Klaim yang dibuktikan", "Artefak", "Tanggal", "Lokasi/tautan", "Status"],
          [
              ["B-01", "Masalah screening tervalidasi praktisi", "Ringkasan wawancara reviewer 1", "[ISI]", "[TAUTAN/hal.]", "[ ]"],
              ["B-06", "Sudut pandang akademik atas metode/etik", "Review akademisi 2", "[ISI]", "[TAUTAN/hal.]", "[ ]"],
              ["D-03", "Warna dashboard direvisi", "Before-after + catatan uji", "[ISI]", "[TAUTAN/hal.]", "[ ]"],
              ["E-02", "CV parsing dan dashboard berfungsi", "Screenshot/demo + test log", "[ISI]", "[TAUTAN/hal.]", "[ ]"],
              ["F-04", "Skor dapat ditelusuri ke evidence", "Contoh scoring trace", "[ISI]", "[TAUTAN/hal.]", "[ ]"],
              ["I-02", "Badan hukum telah terbentuk", "Akta/SK AHU/NIB versi redaksi", "[ISI]", "[TAUTAN/hal.]", "[ ]"],
              ["J-03", "Asumsi BEP 10 perusahaan", "Model finansial dan sensitivitas", "[ISI]", "[TAUTAN/hal.]", "[ ]"],
          ],
          [650, 2300, 2100, 950, 1800, CONTENT_DXA - 7800], font_size=8.2)
add_para(doc, "Template matriks klaim-kriteria", style="Heading 2")
add_table(doc,
          ["Kriteria form", "Kalimat/klaim utama", "Kode bukti", "Apa yang panelis harus lihat"],
          [
              ["Validated User Problem", "[ISI]", "[B-__ / sumber eksternal]", "Profil responden, metode, temuan, tanggal"],
              ["Current Technical Reality", "[ISI]", "[E-__]", "Screenshot, endpoint/test, status aktual"],
              ["Algorithm Quality", "[ISI]", "[F-__]", "Input, bobot/aturan, evidence, edge case, evaluasi"],
              ["User Flow & Testing", "[ISI]", "[D-__]", "Tugas, hasil, masalah, perubahan"],
              ["Team Readiness", "[ISI]", "[I-__]", "Owner, output nyata, legalitas, komitmen"],
              ["Business & ROI", "[ISI]", "[J-__]", "Asumsi, rumus, hasil, sensitivitas"],
          ],
          [1900, 2400, 1500, CONTENT_DXA - 5800], font_size=8.5)

# B validation
add_attachment_heading(
    doc, "B", "Paket Validasi Pengguna dan Expert Review",
    "Mendokumentasikan evidence terbaru dari 4 praktisi industri dan 2 akademisi secara audit-able.",
    "Validated User Problem; Operational Context; Usability Testing; Algorithm Quality; Adoption."
)
add_para(doc, "B.1 Daftar responden/reviewer", style="Heading 2")
add_table(doc,
          ["ID", "Kategori", "Peran/keahlian", "Industri/institusi", "Tanggal", "Metode", "Izin identitas"],
          [
              ["R1", "Industri", "[ISI]", "[ISI]", "[ISI]", "Wawancara / demo / usability", "Nama / anonim"],
              ["R2", "Industri", "[ISI]", "[ISI]", "[ISI]", "[ISI]", "Nama / anonim"],
              ["R3", "Industri", "[ISI]", "[ISI]", "[ISI]", "[ISI]", "Nama / anonim"],
              ["R4", "Industri", "[ISI]", "[ISI]", "[ISI]", "[ISI]", "Nama / anonim"],
              ["R5", "Akademisi", "[ISI]", "[ISI]", "[ISI]", "Expert review / metodologi", "Nama / anonim"],
              ["R6", "Akademisi", "[ISI]", "[ISI]", "[ISI]", "Expert review / etik/AI", "Nama / anonim"],
          ],
          [550, 1000, 1700, 1700, 900, 1750, CONTENT_DXA - 7600], font_size=8.0)
add_callout(doc, "Bukti minimum tiap sesi",
            "Undangan/agenda atau foto (jika diizinkan), catatan/timestamp, daftar pertanyaan atau task, temuan, kutipan singkat, "
            "keputusan tim, dan artefak before-after. Testimoni tanda tangan bersifat pelengkap, bukan bukti tunggal.", AMBER)
add_para(doc, "B.2 Template satu lembar per reviewer (duplikasi 6 kali)", style="Heading 2")
add_template_fields(doc, [
    ("Reviewer ID dan kategori", "R1 - Industri / R5 - Akademisi"),
    ("Profil relevan", "jabatan/keahlian, pengalaman, jenis organisasi; anonimisasi bila perlu"),
    ("Tanggal, durasi, dan metode", "contoh: 20 Juli 2026, 45 menit, moderated prototype walkthrough"),
    ("Tujuan sesi", "klaim/risiko yang hendak divalidasi"),
    ("Tugas atau pertanyaan kunci", "3-5 butir; contoh: review shortlist 20 kandidat dan telusuri alasan skor"),
    ("Temuan utama", "fakta yang diamati, bukan kesimpulan umum"),
    ("Kutipan representatif", "maksimal 1-2 kutipan singkat dengan izin"),
    ("Implikasi terhadap produk/bisnis", "apa yang berubah pada pemahaman masalah atau solusi"),
    ("Keputusan tim", "adopt / modify / reject / backlog, beserta alasan"),
    ("Kode artefak", "tautan rekaman/catatan/screenshot/change log"),
])
add_para(doc, "B.3 Sintesis lintas reviewer", style="Heading 2")
add_table(doc,
          ["Tema", "R1", "R2", "R3", "R4", "R5", "R6", "Kesimpulan tim"],
          [
              ["Beban screening & SLA", "[+/-/NA]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ISI]"],
              ["Kebutuhan explainability", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ISI]"],
              ["Fresh graduate vs profesional", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ISI]"],
              ["Human-in-the-loop", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ISI]"],
              ["Privasi/etik/bias", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ISI]"],
              ["Kesiapan adopsi/integrasi", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ ]", "[ISI]"],
          ],
          [1450, 620, 620, 620, 620, 620, 620, CONTENT_DXA - 5170], font_size=7.8)
add_para(doc, "B.4 Penajaman sejak 2nd submission", style="Heading 2")
add_table(doc,
          ["Pemahaman/fitur sebelumnya", "Evidence baru", "Penajaman", "Perubahan yang dilakukan", "Kode bukti"],
          [
              ["Skor angka sebagai output utama", "[R__]", "Skor perlu kategori + alasan yang dapat diverifikasi", "[ISI]", "[D/E/F-__]"],
              ["Auto-move kandidat ke posisi lain", "[R__]", "Harus menjadi rekomendasi yang disetujui HRD/kandidat", "[ISI]", "[D-__]"],
              ["Warna dashboard dominan", "[R__]", "Mengganggu fokus pada data kandidat", "[ISI]", "[D-__]"],
              ["Alur status belum eksplisit", "[R__]", "Pengguna perlu tahu posisi dan next step", "[ISI]", "[D-__]"],
              ["Input CV dipercaya sebagai data biasa", "[R__]", "Perlu mitigasi prompt injection", "[ISI]", "[E/H-__]"],
          ],
          [1900, 950, 2400, 2500, CONTENT_DXA - 7750], font_size=8.0)

# C problem/usecase
add_attachment_heading(
    doc, "C", "Bukti Problem, Use Case, dan Feature-to-Pain Mapping",
    "Memperlihatkan hubungan masalah -> fitur -> proses -> output -> outcome tanpa lompatan logika.",
    "Validated User Problem; End-to-End Use Case; Operational Context; Complexity."
)
add_para(doc, "C.1 Evidence masalah pengguna", style="Heading 2")
add_table(doc,
          ["Klaim", "Sumber primer", "Sumber sekunder", "Tanggal/versi", "Batasan"],
          [
              ["HRD menghabiskan 5-6 jam/hari", "[ISI: reviewer/data operasional]", "[ISI: publikasi]", "[ISI]", "[ISI]"],
              ["SLA rekrutmen 7-14 hari", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["59% pelamar tidak mendapat feedback", "[ISI]", "Pamungkas (2025) - lengkapi tautan/bibliografi", "[ISI]", "[ISI]"],
              ["22,36% pekerja muda mengalami mismatch", "[ISI]", "Kemnaker RI (2025) - lengkapi dokumen/halaman", "[ISI]", "[ISI]"],
          ],
          [2150, 2200, 2500, 1000, CONTENT_DXA - 7850], font_size=8.2)
add_para(doc, "C.2 Peta end-to-end", style="Heading 2")
add_table(doc,
          ["Tahap", "Kondisi/tindakan", "Input", "Proses sistem", "Output", "Outcome/KPI"],
          [
              ["1. Kondisi awal", "500 CV, hasil interview, tenggat 3 hari", "Lowongan + data kandidat", "-", "Antrian kandidat", "Baseline waktu"],
              ["2. Pemicu", "HRD pilih lowongan dan mulai analisis", "Job criteria", "Validasi job config", "Job siap dianalisis", "Error input rendah"],
              ["3. Screening", "Sistem memproses CV + interview", "CV, audio/transkrip", "Ekstraksi + evidence scoring", "Kategori + alasan", "Waktu screening"],
              ["4. Human review", "HRD menelusuri evidence", "Hasil AI", "Filter + audit trail", "Shortlist/hold/reject", "Agreement rate"],
              ["5. Feedback", "HRD konfirmasi keputusan", "Decision + template", "Generate + review + send", "Email feedback", "Delivery rate"],
              ["6. Outcome", "Proses lebih cepat dan transparan", "Log penggunaan", "Perhitungan KPI", "Laporan pilot", "Time saved, satisfaction"],
          ],
          [1200, 2100, 1450, 1800, 1500, CONTENT_DXA - 8050], font_size=7.8)
add_para(doc, "C.3 Feature-to-pain mapping", style="Heading 2")
add_table(doc,
          ["Pain point", "Fitur", "Mekanisme", "Output yang dilihat", "Outcome & metrik", "Bukti"],
          [
              ["Screening manual lambat", "CV parsing + AI assessment", "Ekstraksi dan pencocokan evidence", "Prioritas kandidat + alasan", "Waktu/kandidat; total waktu", "[E/F-__]"],
              ["Keyword ATS kehilangan konteks", "Evidence-based scoring", "Menilai bukti penggunaan skill", "Matched evidence + gap", "Precision/recall; HR agreement", "[F-__]"],
              ["HRD sulit percaya skor", "Explainable result", "Kutipan CV/transkrip + faktor", "Kategori dan trace", "Trust rating; override rate", "[D/F-__]"],
              ["Kandidat di-ghosting", "AI feedback", "Draft feedback setelah keputusan HRD", "Email personal", "Delivery/read rate", "[D/E-__]"],
          ],
          [1700, 1600, 2050, 1700, 1600, CONTENT_DXA - 8650], font_size=7.7)

# D usability
add_attachment_heading(
    doc, "D", "Usability Testing dan Product Iteration",
    "Membuktikan siapa yang mencoba, tugas yang dilakukan, hasil terukur, masalah, dan perubahan produk.",
    "User Flow, Usability Testing, and Product Iteration; Adoption."
)
add_para(doc, "D.1 Rencana dan sampel pengujian", style="Heading 2")
add_template_fields(doc, [
    ("Tujuan", "contoh: HRD dapat meninjau alasan skor dan membuat shortlist tanpa bantuan"),
    ("Partisipan", "jumlah, profil, pengalaman, kriteria rekrutmen"),
    ("Metode", "moderated/unmoderated, think-aloud, prototype/produk berjalan"),
    ("Perangkat dan versi", "URL/build/commit, desktop/mobile, browser"),
    ("Metrik", "task completion, waktu, error, SEQ/SUS, trust, qualitative notes"),
])
add_para(doc, "D.2 Task sheet", style="Heading 2")
add_table(doc,
          ["Task", "Skenario", "Kriteria berhasil", "Hasil", "Masalah/error", "Evidence"],
          [
              ["T1", "Buka lowongan dan mulai analisis", "[ISI]", "[ISI]", "[ISI]", "[D-01]"],
              ["T2", "Temukan 3 kandidat prioritas dan alasan", "[ISI]", "[ISI]", "[ISI]", "[D-02]"],
              ["T3", "Telusuri evidence dan override bila perlu", "[ISI]", "[ISI]", "[ISI]", "[D-03]"],
              ["T4", "Kirim feedback kandidat tidak lolos", "[ISI]", "[ISI]", "[ISI]", "[D-04]"],
          ],
          [700, 2200, 1950, 1300, 1700, CONTENT_DXA - 7850], font_size=8.1)
add_para(doc, "D.3 Ringkasan hasil", style="Heading 2")
add_table(doc,
          ["Metrik", "Baseline/target", "Hasil aktual", "n", "Interpretasi"],
          [
              ["Task completion", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["Median waktu task", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["Critical error", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["SEQ/SUS/trust score", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
          ],
          [2100, 1700, 1700, 650, CONTENT_DXA - 6150], font_size=8.4)
p = add_para(doc, "D.4 Change log berbasis evidence", style="Heading 2")
p.paragraph_format.page_break_before = True
add_table(doc,
          ["Temuan", "Severity", "Keputusan", "Perubahan", "Before/after", "Status"],
          [
              ["Warna mengalihkan perhatian", "Medium", "Modify", "[ISI]", "[TEMPEL GAMBAR]", "Done"],
              ["Next step/status tidak jelas", "High", "Modify", "[ISI]", "[TEMPEL GAMBAR]", "Done"],
              ["Auto-move mengurangi kontrol", "Critical", "Modify", "[ISI]", "[TEMPEL GAMBAR]", "Done"],
              ["Skor perlu alasan", "High", "Modify", "[ISI]", "[TEMPEL GAMBAR]", "Done"],
          ],
          [1850, 950, 950, 1700, 2300, CONTENT_DXA - 7750], font_size=8.0)
add_callout(doc, "Pencegahan kesalahan",
            "Lampirkan contoh validasi format CV, state loading/progress, konfirmasi sebelum keputusan/kirim feedback, "
            "pesan error yang dapat ditindaklanjuti, retry aman, dan mekanisme HRD mengoreksi hasil AI.", GREEN)

# E technical reality
add_attachment_heading(
    doc, "E", "Bukti Current Technical Reality dan Arsitektur",
    "Membuktikan Innovation Level 3 dan memisahkan komponen aktual dari simulasi/rencana.",
    "Innovation Level; Current Technical Reality; Processing Pipeline; Engineering Depth."
)
add_para(doc, "E.1 Status fitur yang dapat diverifikasi", style="Heading 2")
add_table(doc,
          ["Komponen/fitur", "Status", "Bukti aktual", "Cara verifikasi", "Catatan/batasan"],
          [
              ["Backend REST API", "Berfungsi", "[endpoint/test log]", "[TAUTAN/demo]", "[ISI]"],
              ["PostgreSQL + Redis + object storage", "Berfungsi", "[diagram/config redacted]", "[test]", "[ISI]"],
              ["CV parsing PDF/gambar", "Berfungsi", "[input-output]", "[demo]", "[OCR scan belum]"],
              ["AI assessment + evidence", "Berfungsi", "[screen + JSON]", "[test case]", "[ISI]"],
              ["Transkripsi interview", "Berfungsi", "[screen/log]", "[demo]", "[ISI]"],
              ["Webcam monitoring", "Berfungsi", "[screen]", "[demo]", "[batas pemakaian]"],
              ["Payment subscription", "Simulasi", "[UI/mock]", "-", "[belum transaksi nyata]"],
              ["Chat AI HRD", "Simulasi", "[UI/mock]", "-", "[ISI]"],
              ["Cross-role recommendation", "Dikembangkan", "[branch/design]", "[ISI]", "[belum produksi]"],
              ["Conversational search + OCR scan", "Direncanakan", "[roadmap]", "-", "[ISI]"],
          ],
          [1950, 1100, 1900, 1700, CONTENT_DXA - 6650], font_size=8.0)
add_para(doc, "E.2 Diagram arsitektur", style="Heading 2")
add_callout(doc, "Tempat visual",
            "[TEMPEL: diagram Candidate/HRD Web -> REST API -> PostgreSQL/Redis/Object Storage -> async worker/AI Engine -> external AI providers -> notification/email. "
            "Beri legend untuk data PII, proses asinkron, trust boundary, dan human decision.]", LIGHT_GRAY)
add_para(doc, "E.3 Bukti pipeline input-output", style="Heading 2")
add_table(doc,
          ["Langkah", "Input", "Transformasi/aturan", "Output", "Validasi/fallback", "Evidence"],
          [
              ["Upload", "PDF/image", "Type/size/malware check", "Object reference", "Reject + guidance", "[E-__]"],
              ["Extraction", "CV file", "Parsing/OCR sesuai tipe", "Structured CV", "Missing field flag", "[E-__]"],
              ["Interview", "Audio/video", "Transcription + metadata", "Transcript", "Retry/manual review", "[E-__]"],
              ["Assessment", "CV + transcript + job", "Evidence scoring", "Score/category/reasons", "Schema validation", "[F-__]"],
              ["Dashboard", "Assessment result", "Filter/sort/audit", "HRD view", "Loading/error state", "[E-__]"],
              ["Decision", "HRD action", "Business rules", "Status + feedback", "Confirm + audit log", "[E-__]"],
          ],
          [1150, 1400, 2300, 1550, 1800, CONTENT_DXA - 8200], font_size=7.8)
p = add_para(doc, "E.4 Bukti engineering", style="Heading 2")
p.paragraph_format.page_break_before = True
add_template_fields(doc, [
    ("Reliability", "retry policy, idempotency, caching, timeout, queue, monitoring; sertakan hasil test"),
    ("Performance", "beban uji, ukuran batch, p50/p95 latency, throughput, konfigurasi"),
    ("Scalability", "asynchronous worker, horizontal scaling, rate limiting, cost guardrail"),
    ("Modularity", "kontrak API, pemisahan web/API/AI engine, model-provider abstraction"),
    ("Deployment", "diagram environment, versi, tanggal deploy, health check, rollback/backup"),
])

# F algorithm
add_attachment_heading(
    doc, "F", "Algorithm/Rule Quality dan Decision Transparency",
    "Membuat Evidence-Based Scoring dapat dipahami, diuji, ditelusuri, dan dikoreksi.",
    "Algorithm or Rule Quality and Decision Transparency; Complexity."
)
add_para(doc, "F.1 Spesifikasi scoring", style="Heading 2")
add_table(doc,
          ["Aspek", "Input/evidence", "Aturan/bobot", "Output", "Pengecualian"],
          [
              ["Keahlian", "Skill + konteks penggunaan", "[ISI bobot/rubrik]", "Matched/partial/missing", "Skill hanya disebut"],
              ["Pengalaman", "Peran, durasi, relevansi", "[ISI]", "Evidence + score", "Tanggal ambigu/gap"],
              ["Pendidikan", "Jenjang/bidang/sertifikasi", "[ISI]", "Match/related/not found", "Equivalent credential"],
              ["Peran & tanggung jawab", "Scope dan aktivitas", "[ISI]", "Evidence + strength", "Judul jabatan berbeda"],
              ["Nilai tambah", "Proyek, organisasi, prestasi", "[ISI]", "Bonus/cap", "Fresh graduate bias"],
          ],
          [1450, 2100, 1750, 1700, CONTENT_DXA - 7000], font_size=8.1)
add_para(doc, "F.2 Contoh trace satu keputusan", style="Heading 2")
add_table(doc,
          ["Requirement", "Evidence kandidat", "Sumber/lokasi", "Penilaian", "Alasan", "HRD correction"],
          [
              ["[ISI]", "[kutipan singkat]", "CV p.__ / interview 00:__", "[match/partial/missing]", "[ISI]", "[accept/edit/reject]"],
              ["[ISI]", "[ISI]", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["[ISI]", "[ISI]", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
          ],
          [1600, 2150, 1350, 1300, 1900, CONTENT_DXA - 8300], font_size=7.8)
add_para(doc, "F.3 Evaluation set dan quality metrics", style="Heading 2")
add_table(doc,
          ["Dimensi", "Definisi", "Metode", "Baseline", "Hasil", "Ambang lulus"],
          [
              ["Extraction accuracy", "[ISI]", "Ground truth vs output", "[ISI]", "[ISI]", "[ISI]"],
              ["Requirement classification", "Match/partial/missing", "Labeled test set", "[keyword baseline]", "[ISI]", "[ISI]"],
              ["Ranking agreement", "Keselarasan urutan AI-HRD", "Top-k/Spearman/agreement", "[ISI]", "[ISI]", "[ISI]"],
              ["Evidence faithfulness", "Alasan didukung sumber", "Reviewer audit", "[ISI]", "[ISI]", "[ISI]"],
              ["Consistency", "Output stabil untuk input sama", "Repeat test", "[ISI]", "[ISI]", "[ISI]"],
              ["Fairness slice", "Performa antar tipe kandidat", "Fresh grad vs profesional", "[ISI]", "[ISI]", "[ISI]"],
          ],
          [1700, 2000, 1900, 1300, 1300, CONTENT_DXA - 8200], font_size=7.8)
p = add_para(doc, "F.4 Edge cases dan guardrails", style="Heading 2")
p.paragraph_format.page_break_before = True
add_table(doc,
          ["Kasus", "Risiko", "Perilaku sistem yang diharapkan", "Test result", "Owner"],
          [
              ["CV kosong/rusak", "False assessment", "Stop, minta unggah ulang", "[ISI]", "Engineering"],
              ["Prompt injection dalam CV", "Manipulasi hasil", "Treat as data, structured output, filter", "[ISI]", "AI/Security"],
              ["Informasi bertentangan", "Alasan tidak konsisten", "Flag untuk HRD", "[ISI]", "AI/Data"],
              ["Evidence tidak cukup", "Overconfidence", "Abstain/perlu review", "[ISI]", "Product/HRD"],
              ["Kandidat meminta koreksi", "Data tidak akurat", "Correction workflow + audit", "[ISI]", "Operations"],
          ],
          [1750, 1800, 3000, 1300, CONTENT_DXA - 7850], font_size=8.0)

# G MVP
add_attachment_heading(
    doc, "G", "MVP Execution, Deployment, dan Risk Register",
    "Menunjukkan scope pilot, milestone, PIC, kebutuhan integrasi, operasi, dan mitigasi yang konkret.",
    "MVP Execution and Deployment Plan; Continuation Readiness."
)
add_para(doc, "G.1 Scope MVP/pilot", style="Heading 2")
add_table(doc,
          ["In scope", "Acceptance criteria", "Out of scope", "Alasan"],
          [
              ["Create job, apply, CV parsing", "[ISI]", "Auto-transfer kandidat", "Perlu consent + HRD approval"],
              ["Pre-screening + AI interview", "[ISI]", "Conversational candidate search", "Fokus core selection"],
              ["Evidence assessment + HRD dashboard", "[ISI]", "Real payment processing", "Belum dibutuhkan pilot"],
              ["Decision + automated feedback draft", "[ISI]", "OCR dokumen scan lanjutan", "Tahap berikutnya"],
          ],
          [2300, 2500, 2300, CONTENT_DXA - 7100], font_size=8.2)
add_para(doc, "G.2 Milestone 6-12 bulan", style="Heading 2")
add_table(doc,
          ["Waktu", "Milestone", "Output/DoD", "PIC", "Sumber daya", "Dependency"],
          [
              ["Bulan 1", "Pilot readiness", "[security/privacy + test + onboarding]", "Raffi/Randi", "[ISI]", "Mitra pilot"],
              ["Bulan 2-3", "1 pilot terbatas", "[volume, role, KPI, report]", "All; owner Raffi", "[ISI]", "Data & user"],
              ["Bulan 4-6", "Product hardening", "[reliability + scoring eval]", "Alief/Jan", "[ISI]", "AI provider"],
              ["Bulan 7-9", "Expansion within partner", "[divisi/role tambahan]", "Raffi/Randi", "[ISI]", "Pilot success"],
              ["Bulan 10-12", "Paid conversion", "[pricing test + customers]", "Raffi", "[ISI]", "Legal/billing"],
          ],
          [1100, 1900, 2450, 1400, 1300, CONTENT_DXA - 8150], font_size=7.9)
add_para(doc, "G.3 Risk register", style="Heading 2")
add_table(doc,
          ["Risiko", "Kategori", "Prob.", "Dampak", "Mitigasi", "Trigger/contingency", "Owner"],
          [
              ["Rate limit/API outage", "Teknis", "M", "H", "Retry, queue, cache, provider abstraction", "[ISI]", "Alief"],
              ["AI result tidak konsisten", "Teknis", "M", "H", "Schema, eval set, human review", "[ISI]", "Jan"],
              ["Kebocoran PII", "Legal/teknis", "L/M", "Critical", "Access control, retention, incident plan", "[ISI]", "Randi"],
              ["Pilot tidak menyediakan data", "Operasional", "M", "H", "Synthetic/demo set + DPA + minimal scope", "[ISI]", "Raffi"],
              ["Biaya AI melebihi asumsi", "Finansial", "M", "M/H", "Quota, cost logging, model routing", "[ISI]", "Raffi/Jan"],
          ],
          [1800, 1100, 650, 800, 2450, 1800, CONTENT_DXA - 8600], font_size=7.7)
p = add_para(doc, "G.4 Runbook operasi", style="Heading 2")
p.paragraph_format.page_break_before = True
add_template_fields(doc, [
    ("Environment & deployment", "Vercel/Heroku atau kondisi terbaru, branch/release policy, secrets"),
    ("Monitoring", "uptime, error rate, queue depth, AI latency/cost, alert owner"),
    ("Backup & restore", "cakupan, frekuensi, retention, bukti restore test"),
    ("Incident response", "severity, kontak, response time, notification, postmortem"),
    ("Maintenance", "jadwal, dependency updates, model/version change control"),
])

# H Privacy legal
add_attachment_heading(
    doc, "H", "Keamanan, Privasi, Kepatuhan, dan Solution Boundary",
    "Menunjukkan perlindungan PII serta batas keputusan AI secara operasional.",
    "Operational Context; Current Technical Reality; Adoption; Legal Risk."
)
add_para(doc, "H.1 Data inventory", style="Heading 2")
add_table(doc,
          ["Data", "Sumber", "Tujuan", "Lokasi", "Akses", "Retention/deletion", "Dasar/persetujuan"],
          [
              ["Identitas & kontak", "Kandidat", "Aplikasi & komunikasi", "[ISI]", "[role]", "[ISI]", "[ISI]"],
              ["CV & riwayat", "Kandidat", "Screening", "[ISI]", "[role]", "[ISI]", "[ISI]"],
              ["Audio/video/transkrip", "Interview", "Assessment/monitoring", "[ISI]", "[role]", "[ISI]", "[ISI]"],
              ["Job criteria", "Perusahaan", "Scoring", "[ISI]", "[role]", "[ISI]", "[ISI]"],
              ["Assessment & feedback", "Sistem/HRD", "Decision support", "[ISI]", "[role]", "[ISI]", "[ISI]"],
          ],
          [1650, 1200, 1700, 1300, 1000, 1650, CONTENT_DXA - 8500], font_size=7.6)
add_para(doc, "H.2 Kontrol dan bukti", style="Heading 2")
add_table(doc,
          ["Area", "Kontrol", "Status", "Bukti", "Gap & tanggal penyelesaian"],
          [
              ["Access", "Role-based access, least privilege", "[ISI]", "[H-__]", "[ISI]"],
              ["Encryption", "In transit/at rest sesuai layanan", "[ISI]", "[H-__]", "[ISI]"],
              ["Secrets", "Tidak disimpan di repo/log", "[ISI]", "[H-__]", "[ISI]"],
              ["Logging", "Audit access/decision tanpa PII berlebih", "[ISI]", "[H-__]", "[ISI]"],
              ["Retention", "Delete/export request & policy", "[ISI]", "[H-__]", "[ISI]"],
              ["AI provider", "DPA/terms/data use review", "[ISI]", "[H-__]", "[ISI]"],
              ["Prompt injection", "Input isolation + validation + tests", "[ISI]", "[H-__]", "[ISI]"],
          ],
          [1450, 3000, 1100, 1200, CONTENT_DXA - 6750], font_size=8.0)
add_para(doc, "H.3 Pernyataan batas solusi", style="Heading 2")
add_template_fields(doc, [
    ("Yang dapat dilakukan", "ekstraksi, interview, decision support, feedback draft, dashboard"),
    ("Yang tidak dilakukan", "tidak membuat keputusan final otomatis; tidak menjamin kecocokan kerja"),
    ("Human decision point", "siapa memvalidasi, kapan, data apa yang dilihat, bagaimana override dicatat"),
    ("Hak kandidat", "informasi penggunaan AI, consent, correction, withdrawal/deletion sesuai kebijakan"),
    ("Ketergantungan eksternal", "AI provider, email, cloud, regulator, mitra pelatihan, ATS/HRIS"),
])
add_callout(doc, "Dokumen legal",
            "Lampirkan Privacy Notice, consent notice untuk AI interview/rekaman, Terms, data processing checklist/DPA pilot, "
            "dan ringkasan kepatuhan UU PDP yang telah direview. Hindari klaim 'sudah patuh' tanpa review yang dapat dibuktikan.", AMBER)

# I team legal
add_attachment_heading(
    doc, "I", "Team Capability, Ownership, dan Legal Entity",
    "Membuktikan siapa mengerjakan apa, output aktual, komitmen, dan kesiapan organisasi.",
    "Team Capability and Execution Ownership; Continuation Readiness."
)
add_para(doc, "I.1 Ownership dan hasil kerja", style="Heading 2")
add_table(doc,
          ["Anggota", "Peran", "Ownership", "Hasil nyata", "Kompetensi/evidence", "Milestone berikut"],
          [
              ["Raffi Rabbani Widyputra", "Product & Business Lead", "Roadmap, discovery, pilot, pricing", "[ISI + link]", "[ISI]", "Pilot & conversion"],
              ["Alief Athallah Putra", "Head of Engineering", "Web/API/infra/reliability", "[ISI + link]", "[ISI]", "Production hardening"],
              ["Jan Agra Adyuta Harnowo", "AI & Data Lead", "Scoring, eval, AI integration", "[ISI + link]", "[ISI]", "Evaluation benchmark"],
              ["Randi Adam Arnaldi", "Operations & Compliance Lead", "Ops, privacy, legal, partnership", "[ISI + link]", "[ISI]", "Pilot ops & DPA"],
          ],
          [1700, 1600, 2200, 1700, 1500, CONTENT_DXA - 8700], font_size=7.7)
add_para(doc, "I.2 Cara mengambil keputusan", style="Heading 2")
add_template_fields(doc, [
    ("Forum dan ritme", "contoh: weekly product review; monthly business review"),
    ("Decision owner", "Raffi untuk prioritas produk; domain owner untuk keputusan teknis/AI/compliance"),
    ("Dokumentasi", "decision log, backlog, metric review, dissent/escalation"),
    ("Komitmen", "jam/orang/minggu, biaya/cloud budget, periode 6-12 bulan"),
    ("Gap kompetensi", "HR advisor, legal PDP, enterprise sales, ML evaluation; cara mendapatkannya"),
])
add_para(doc, "I.3 Checklist bukti badan hukum", style="Heading 2")
add_table(doc,
          ["Dokumen", "Lampirkan?", "Versi publik/redaksi", "Catatan"],
          [
              ["Akta pendirian", "[ ]", "Halaman identitas badan + maksud usaha; redaksi data pribadi", "[ISI]"],
              ["SK pengesahan AHU", "[ ]", "Nomor dapat disamarkan sebagian bila perlu", "[ISI]"],
              ["NIB/izin berbasis risiko", "[ ]", "Tampilkan KBLI relevan; redaksi QR/nomor sensitif", "[ISI]"],
              ["NPWP badan", "Opsional", "Jangan tampilkan nomor penuh di paket publik", "[ISI]"],
              ["Struktur kepemilikan/pendiri", "[ ]", "Ringkas; konsisten dengan tim", "[ISI]"],
              ["Perjanjian founders/IP assignment", "Jika ada", "Ringkasan status, bukan isi rahasia penuh", "[ISI]"],
          ],
          [1900, 1000, 4150, CONTENT_DXA - 7050], font_size=8.1)
add_callout(doc, "Redaksi wajib",
            "Hapus atau tutup NIK, NPWP pribadi, alamat rumah, tanda tangan, nomor rekening, QR code, email/telepon pribadi, "
            "dan detail beneficial ownership yang tidak diperlukan untuk penilaian.", RED)

# J financial
add_attachment_heading(
    doc, "J", "Financial Analysis, Business Model, dan ROI",
    "Menunjukkan bahwa pricing, biaya, nilai pelanggan, BEP, dan sensitivitas berasal dari asumsi yang dapat ditelusuri.",
    "Quantified Value, Business Model, and ROI."
)
add_para(doc, "J.1 Actor-value-money map", style="Heading 2")
add_table(doc,
          ["Pihak", "Menggunakan", "Menerima manfaat", "Membayar/mendanai", "Unit manfaat/KPI"],
          [
              ["HRD/recruiter", "Dashboard & assessment", "Waktu screening turun", "Perusahaan", "Jam hemat, throughput"],
              ["Hiring manager", "Shortlist/evidence", "Keputusan lebih cepat", "Perusahaan", "Time-to-shortlist"],
              ["Kandidat", "Apply/interview/feedback", "Kepastian + pengembangan", "Tidak/opsional", "Feedback delivery"],
              ["Perusahaan", "Platform", "Produktivitas + risk reduction", "Subscription/API/posting", "ROI/payback"],
              ["Mitra pelatihan/pemerintah", "Referral insight", "Peserta relevan/impact", "[ISI]", "Referral/completion"],
          ],
          [1500, 1700, 2100, 1900, CONTENT_DXA - 7200], font_size=8.1)
add_para(doc, "J.2 Asumsi inti dan sumber", style="Heading 2")
add_table(doc,
          ["Asumsi", "Nilai", "Satuan", "Sumber/evidence", "Confidence", "Sensitivity"],
          [
              ["Gaji HRD", "Rp8.000.000", "per bulan", "[sumber/benchmark]", "M", "[low/base/high]"],
              ["Waktu screening", "5-6", "jam/hari", "[review/data]", "M", "[ISI]"],
              ["Hari kerja", "[ISI]", "hari/bulan", "[ISI]", "H", "[ISI]"],
              ["Penghematan waktu", "[ISI]", "% atau jam", "[pilot/benchmark]", "L/M", "[ISI]"],
              ["Harga Pro", "Rp299-500 ribu", "per bulan", "[pricing test]", "L/M", "[ISI]"],
              ["AI cost/session", "[ISI]", "Rp/interview", "[provider calculator/log]", "M", "[ISI]"],
              ["Conversion free-paid", "20%", "%", "[hipotesis/benchmark]", "L", "[10/20/30%]"],
              ["BEP customer", "10", "perusahaan", "[model]", "L/M", "[ISI]"],
          ],
          [1900, 1400, 1300, 2300, 950, CONTENT_DXA - 7850], font_size=8.0)
p = add_para(doc, "J.3 Formula ROI pelanggan", style="Heading 2")
p.paragraph_format.page_break_before = True
add_callout(doc, "Rumus",
            "Nilai waktu hemat = jam screening baseline x % penghematan x biaya tenaga kerja per jam. "
            "ROI = (nilai manfaat tahunan - biaya tahunan Direkrut AI - biaya implementasi) / total biaya. "
            "Payback period = total biaya implementasi / manfaat bersih bulanan.", GREEN)
add_table(doc,
          ["Skenario", "Low", "Base", "High", "Catatan"],
          [
              ["Jumlah recruiter", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["Lamaran/bulan", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["Jam hemat/bulan", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["Manfaat Rp/bulan", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["Biaya Rp/bulan", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["ROI tahunan", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["Payback", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
          ],
          [2400, 1400, 1400, 1400, CONTENT_DXA - 6600], font_size=8.3)
p = add_para(doc, "J.4 Unit economics & 12-month summary", style="Heading 2")
p.paragraph_format.page_break_before = True
add_table(doc,
          ["Metrik", "Formula", "M1", "M6", "M12", "Target/interpretasi"],
          [
              ["MRR", "Paid customer x ARPA", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["COGS", "AI + voice + cloud + support variable", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["Gross margin", "(Revenue - COGS)/Revenue", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["CAC", "Sales & marketing/new customer", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["LTV", "ARPA x GM / churn", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["Cash need", "Opex + capex - receipts", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
          ],
          [1600, 2500, 1000, 1000, 1000, CONTENT_DXA - 7100], font_size=8.0)
add_callout(doc, "Lampiran spreadsheet",
            "Sertakan file financial model dengan tab Assumptions, Pricing, Unit Economics, 12M P&L/Cashflow, ROI Customer, "
            "dan Sensitivity. Tandai angka proposal sebagai base case; bedakan benchmark, hipotesis, dan hasil pilot.", AMBER)

# K growth
add_attachment_heading(
    doc, "K", "Adoption, Pilot, Growth, dan Competitive Moat",
    "Membuktikan jalur memperoleh mitra pertama dan menguji keunggulan yang sulit digantikan.",
    "Adoption, Growth Strategy, and Competitive Moat; Continuation Readiness."
)
add_para(doc, "K.1 Pilot funnel", style="Heading 2")
add_table(doc,
          ["Target organisasi", "Kontak/peran", "Problem fit", "Tahap", "Next step/tanggal", "Evidence"],
          [
              ["[ISI]", "[ISI]", "[ISI]", "Intro/discovery/demo/LOI/pilot", "[ISI]", "[email/notes]"],
              ["[ISI]", "[ISI]", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
              ["[ISI]", "[ISI]", "[ISI]", "[ISI]", "[ISI]", "[ISI]"],
          ],
          [1800, 1500, 1900, 1700, 1700, CONTENT_DXA - 8600], font_size=8.0)
add_para(doc, "K.2 Pilot one-pager", style="Heading 2")
add_template_fields(doc, [
    ("Mitra dan use case", "jenis organisasi, role yang direkrut, volume kandidat"),
    ("Periode dan scope", "tanggal, fitur, jumlah recruiter/kandidat"),
    ("Success metrics", "time-to-shortlist, time saved, agreement, feedback delivery, satisfaction"),
    ("Data & legal", "data minimum, consent, DPA, retention, akses"),
    ("Dukungan", "onboarding, SLA, escalation, training"),
    ("Exit/scale criteria", "kapan stop, extend, atau convert berbayar"),
])
add_para(doc, "K.3 Competitive moat hypothesis", style="Heading 2")
add_table(doc,
          ["Moat candidate", "Kenapa bernilai", "Bukti sekarang", "Cara diperkuat", "Risiko mudah ditiru"],
          [
              ["Evidence trace & HR feedback loop", "Trust + proprietary workflow data", "[ISI]", "Correction data + eval set", "[ISI]"],
              ["Rubrik fresh grad vs profesional", "Lebih relevan per segmen", "[ISI]", "Domain calibration", "[ISI]"],
              ["End-to-end candidate feedback", "Experience + completion loop", "[ISI]", "Partner training network", "[ISI]"],
              ["Integration/workflow switching cost", "Tertanam di proses HR", "[ISI]", "ATS/HRIS integration", "[ISI]"],
          ],
          [1950, 2100, 1750, 2000, CONTENT_DXA - 7800], font_size=7.9)
add_callout(doc, "Bukti ketertarikan yang kuat",
            "Urutan kekuatan bukti: pilot berjalan > LOI/MoU > email komitmen dengan scope/tanggal > demo terjadwal > catatan discovery > testimoni umum. "
            "Jangan menyebut 'mitra' bila baru sebatas reviewer tanpa pembicaraan pilot.", AMBER)

# L pitch
add_attachment_heading(
    doc, "L", "Pitch Deck, Demo, dan Referensi",
    "Mengemas ringkasan visual tanpa menggantikan evidence rinci.",
    "Seluruh kriteria, terutama clarity, traction, feasibility, dan business."
)
add_para(doc, "L.1 Checklist pitch deck", style="Heading 2")
for item in [
    "Problem dan bukti terbaru, termasuk insight 4 praktisi + 2 akademisi.",
    "Satu end-to-end use case dengan masalah -> fitur -> output -> outcome.",
    "Demo/screenshot produk aktual dengan label berfungsi/simulasi.",
    "Evidence-Based Scoring, explainability, human validation, dan guardrail.",
    "Hasil usability/review dan before-after perubahan sejak 2nd submission.",
    "Business model, ROI base case, unit economics, dan sensitivity singkat.",
    "Pilot plan, traction/LOI bila ada, roadmap, team ownership, dan legal entity.",
    "Ask yang spesifik: mitra pilot, akses data, advisor HR/legal/AI, atau pendanaan.",
]:
    add_bullet(doc, item)
add_para(doc, "L.2 Demo evidence", style="Heading 2")
add_table(doc,
          ["Scene", "Apa yang ditampilkan", "Durasi", "Klaim", "Bukti cadangan"],
          [
              ["1", "Create/select job + criteria", "[ISI]", "Context-aware setup", "[screenshot]"],
              ["2", "CV upload/extraction", "[ISI]", "Structured input", "[input-output]"],
              ["3", "AI interview/transcript", "[ISI]", "Multimodal flow", "[screen recording]"],
              ["4", "Assessment/evidence trace", "[ISI]", "Transparent decision support", "[JSON/test]"],
              ["5", "HRD decision + feedback", "[ISI]", "Human control + anti-ghosting", "[email preview]"],
          ],
          [650, 2700, 900, 2300, CONTENT_DXA - 6550], font_size=8.2)
add_para(doc, "L.3 Daftar referensi", style="Heading 2")
add_table(doc,
          ["Ref", "Sitasi lengkap", "Tautan/DOI", "Tanggal akses", "Klaim yang didukung"],
          [
              ["S1", "Pamungkas (2025) - [lengkapi judul, penerbit]", "[ISI]", "[ISI]", "59% pelamar tanpa feedback"],
              ["S2", "Kemnaker RI (2025) - [lengkapi dokumen/tabel]", "[ISI]", "[ISI]", "22,36% mismatch"],
              ["S3", "Suhendra & Supriatin (2025) - [lengkapi]", "[ISI]", "[ISI]", "Time/cost improvement"],
              ["S4", "Harvard Business Review - [verifikasi sumber tepat]", "[ISI]", "[ISI]", "Cost of bad hire"],
              ["S5", "[Sumber regulasi/UU PDP]", "[ISI]", "[ISI]", "Privacy & compliance"],
          ],
          [650, 3300, 1800, 1200, CONTENT_DXA - 6950], font_size=8.0)
add_callout(doc, "Verifikasi referensi",
            "Lengkapi judul, penerbit, URL, tanggal, dan halaman/tabel. Jangan menggunakan klaim sekunder yang tidak dapat dibuka panelis. "
            "Untuk angka ROI, tandai apakah sumber adalah benchmark eksternal atau hasil Direkrut AI.", RED)

# Final checklist
doc.add_page_break()
add_para(doc, "FINAL QUALITY GATE", bold=True, color=TEAL, size=10, after=3)
add_para(doc, "Checklist sebelum dikirim", style="Heading 1")
checks = [
    "Semua placeholder [ISI], [TAUTAN], dan [TEMPEL] telah diisi atau dihapus.",
    "Seluruh angka pada form memiliki sumber atau asumsi di Lampiran C/J/L.",
    "Enam reviewer memiliki catatan individual dan sintesis lintas reviewer.",
    "Nama/identitas reviewer hanya ditampilkan sesuai izin; transkrip dan rekaman memiliki consent.",
    "Setiap screenshot diberi tanggal/versi, caption, dan label status aktual.",
    "Demo tidak mencampur fitur berfungsi dengan simulasi tanpa label.",
    "Scoring memiliki satu contoh trace end-to-end dan hasil quality test.",
    "Human-in-the-loop, exception, override, dan correction path terlihat.",
    "Dokumen legal dan screenshot sistem telah disunting dari data sensitif.",
    "Financial model membedakan benchmark, hipotesis, dan hasil aktual.",
    "Pitch deck konsisten dengan proposal, lampiran, demo, pricing, dan team role.",
    "Tautan menggunakan akses read-only dan sudah diuji dari akun lain/incognito.",
    "Nama file, nomor versi, tanggal, dan indeks halaman konsisten.",
]
for item in checks:
    add_para(doc, f"[ ] {item}", after=4)
add_para(doc, "Konvensi nama file", style="Heading 2")
add_callout(doc, "Format",
            "S0293_DirekrutAI_[Kode]_[NamaDokumen]_v[Versi]_[YYYYMMDD].pdf\n"
            "Contoh: S0293_DirekrutAI_B_ValidasiReviewer_v1_20260726.pdf", LIGHT_GRAY)
add_para(doc, "Struktur folder yang disarankan", style="Heading 2")
add_table(doc,
          ["Folder", "Isi"],
          [
              ["00_Evidence_Index", "Master index dan cross-reference"],
              ["01_User_Validation", "Consent, notes, synthesis, external sources"],
              ["02_Product_UX", "Flow, screenshots, usability, change log"],
              ["03_Technical_Algorithm", "Architecture, test, scoring, security"],
              ["04_MVP_Operations", "Roadmap, risk, runbook, pilot"],
              ["05_Team_Legal", "Profiles, role evidence, legal entity redacted"],
              ["06_Business_Growth", "Financial model, ROI, LOI, pitch deck"],
          ],
          [2700, CONTENT_DXA - 2700], font_size=8.8)
add_para(doc, "Persetujuan final", style="Heading 2")
add_table(doc,
          ["Area", "Reviewer internal", "Tanggal", "Status/catatan"],
          [
              ["Product & user evidence", "Raffi", "[ISI]", "[ ]"],
              ["Technical & deployment", "Alief", "[ISI]", "[ ]"],
              ["AI quality & data", "Jan", "[ISI]", "[ ]"],
              ["Operations, legal & privacy", "Randi", "[ISI]", "[ ]"],
              ["Cross-document consistency", "Seluruh tim", "[ISI]", "[ ]"],
          ],
          [2400, 1700, 1300, CONTENT_DXA - 5400], font_size=8.5)

# Metadata
doc.core_properties.title = "Template Lampiran Proposal Direkrut AI"
doc.core_properties.subject = "Evidence Pack Digdaya X Hackathon"
doc.core_properties.author = "Team Direkrut AI"
doc.core_properties.keywords = "Direkrut AI, Hackathon, Lampiran, Evidence Pack"

doc.save(OUT)
print(OUT)
